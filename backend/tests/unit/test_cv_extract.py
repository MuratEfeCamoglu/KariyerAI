import io

import docx
import pytest

from app.services.cv_extract import ExtractionError, FileKind, detect_kind, extract_text
from tests.conftest import FIXTURES

PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\x00" * 64


def test_detects_pdf_and_docx_from_content() -> None:
    assert detect_kind((FIXTURES / "cv_deniz_arslan.pdf").read_bytes()) is FileKind.PDF
    assert detect_kind((FIXTURES / "cv_elif_kaya.docx").read_bytes()) is FileKind.DOCX


@pytest.mark.parametrize(
    "data",
    [PNG_BYTES, b"Duz metin dosyasi, uzantisi .pdf olsa bile", b"", b"PK\x03\x04bozuk-zip"],
)
def test_rejects_other_content(data: bytes) -> None:
    assert detect_kind(data) is None


def test_extracts_pdf_text() -> None:
    text = extract_text((FIXTURES / "cv_deniz_arslan.pdf").read_bytes(), FileKind.PDF)

    assert text.startswith("Deniz Arslan")
    assert "Skills: Python, FastAPI, PostgreSQL, Docker" in text
    assert len(text) >= 200


def test_extracts_docx_text_with_turkish_characters() -> None:
    text = extract_text((FIXTURES / "cv_elif_kaya.docx").read_bytes(), FileKind.DOCX)

    assert "Beceriler: SQL, Python, Power BI, Excel" in text
    assert "İngilizce (B2)" in text


def test_extracts_docx_tables() -> None:
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Beceri"
    table.cell(0, 1).text = "Kubernetes"
    buffer = io.BytesIO()
    document.save(buffer)

    assert "Beceri | Kubernetes" in extract_text(buffer.getvalue(), FileKind.DOCX)


def test_image_only_pdf_yields_no_text() -> None:
    assert extract_text((FIXTURES / "scanned_cv.pdf").read_bytes(), FileKind.PDF) == ""


def test_corrupt_files_raise_extraction_error() -> None:
    with pytest.raises(ExtractionError):
        extract_text(b"%PDF-1.4\nbu bir pdf degil", FileKind.PDF)
    with pytest.raises(ExtractionError):
        extract_text(b"PK\x03\x04bozuk", FileKind.DOCX)
