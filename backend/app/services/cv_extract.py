import io
import re
import zipfile
from enum import StrEnum

import docx
import filetype
from pypdf import PdfReader
from pypdf.errors import PdfReadError

MAX_FILE_BYTES = 5 * 1024 * 1024
MIN_TEXT_LENGTH = 200

_MIME_TO_KIND = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}


class FileKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"


class ExtractionError(Exception):
    """Dosya açılamadı veya okunamadı (bozuk, şifreli vb.)."""


def detect_kind(data: bytes) -> FileKind | None:
    """Türü içerikten (magic bytes) belirler; dosya uzantısına güvenilmez (F4c)."""
    guess = filetype.guess(data)
    kind = _MIME_TO_KIND.get(guess.mime) if guess is not None else None
    return FileKind(kind) if kind is not None else None


def _normalize(text: str) -> str:
    lines = (re.sub(r"[ \t ]+", " ", line).strip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ExtractionError("Şifreli PDF okunamadı")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    # CV'lerde beceri/deneyim sık sık tablo içinde olur.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(data: bytes, kind: FileKind) -> str:
    try:
        raw = _extract_pdf(data) if kind is FileKind.PDF else _extract_docx(data)
    except ExtractionError:
        raise
    except (PdfReadError, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ExtractionError("Dosya okunamadı") from exc
    return _normalize(raw)
